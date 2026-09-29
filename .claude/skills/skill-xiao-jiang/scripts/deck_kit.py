# -*- coding: utf-8 -*-
"""小講的講座簡報工具箱 — AI數字員工龍蝦學院品牌模組 + helper。

用法：在你的 build_ppt.py 裡
    from deck_kit import *
    prs = new_deck()
    s = slide(prs)            # 一頁黑底 + 右上 logo
    box(s, ...)               # 加多行文字
    link_box(s, ...)          # 加可點超連結
    add_countdown_slide(prs, COUNTDOWN_MP4, POSTER_PNG)  # 倒數待機頁(自動播)
    prs.save(OUT)
然後另跑 add_notes 把逐字稿灌進 speaker notes（重建會洗掉 notes，務必後跑）。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
import os

# ── 品牌（本機規則：預設不掛）─────────────────────────────
# 對外產出預設不掛任何品牌。要掛龍蝦學院品牌時：
#     export LOBSTER_BRAND="AI數字員工龍蝦學院"
# 判斷準則見 skill-ai-lobster-brand 的「適用範圍」。
BRAND = os.environ.get("LOBSTER_BRAND", "")

# ---- 品牌色（黑底 + 金 + 龍蝦橘）----
INK    = RGBColor(0x12, 0x12, 0x16)
PANEL  = RGBColor(0x1E, 0x1E, 0x26)
GOLD   = RGBColor(0xE8, 0xB7, 0x4A)
ORANGE = RGBColor(0xF2, 0x6B, 0x2C)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
MUTE   = RGBColor(0xB8, 0xBD, 0xC7)
LIGHT  = RGBColor(0xF7, 0xF6, 0xF2)

CN = "Microsoft JhengHei"
EW = Inches(13.333)
EH = Inches(7.5)
LOGO = "/Users/chenyunung/.claude/skills/skill-ai-lobster-brand/logo.png"

# 影片自動播放 timing：{spid} 填影片 shape id；放映載入該頁即自動播。
AUTOPLAY_TIMING = (
    '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
    '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
    '<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
    '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
    '<p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '<p:childTnLst><p:par>'
    '<p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '<p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr>'
    '<p:cTn id="6" dur="1801000" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    '</p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst>'
    '</p:cTn></p:par></p:childTnLst></p:cTn>'
    '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
    '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
    '</p:seq>'
    '<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst>'
    '</p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>'
    '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
)


def new_deck():
    prs = Presentation()
    prs.slide_width = EW
    prs.slide_height = EH
    return prs


def slide(prs, bg=INK):
    """新增一頁：底色 + 右上角 logo。"""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EW, EH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    s.shapes._spTree.remove(r._element)
    s.shapes._spTree.insert(2, r._element)
    s.shapes.add_picture(LOGO, Inches(12.55), Inches(0.15),
                         width=Inches(0.6), height=Inches(0.6))
    return s


def box(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines = [(text, size_pt, color, bold), ...]"""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, sz, col, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(6)
        r = p.add_run(); r.text = txt
        f = r.font; f.size = Pt(sz); f.color.rgb = col; f.bold = bold; f.name = CN
    return tb


def bar(s, x, y, w, h, col):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = col
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp


def panel(s, x, y, w, h, col=PANEL):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = col
    sp.line.color.rgb = GOLD; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp


def kicker(s, txt, y=Inches(0.9), col=GOLD):
    box(s, Inches(0.9), y, Inches(10), Inches(0.5), [(txt, 16, col, True)])


def link_box(s, x, y, w, label, url, col=GOLD, sz=18):
    """帶 https 超連結的一行字（單擊/Cmd+click 觸發；本機 file:// 在 Mac 沙盒 PPT 會被擋）。"""
    tb = s.shapes.add_textbox(x, y, w, Inches(0.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = label
    f = r.font; f.size = Pt(sz); f.color.rgb = col; f.bold = True; f.name = CN
    r.hyperlink.address = url
    return tb


def add_countdown_slide(prs, mp4_path, poster_path,
                        title="直播即將開始",
                        subtitle="用 AI 養出一池會幫你賺錢的「數字員工」",
                        footer="先在留言區打個「1」，讓我知道你在線上 ▎ 我們 30 分鐘後見！",
                        sign=""):
    """直播倒數待機頁：金框倒數影片 + 自動播放。回傳 slide。"""
    s = slide(prs, INK)
    box(s, Inches(0), Inches(0.55), EW, Inches(0.5),
        [((f"{BRAND} ▎ LIVE" if BRAND else "LIVE"), 18, GOLD, True)], align=PP_ALIGN.CENTER)
    box(s, Inches(0), Inches(1.02), EW, Inches(1.0),
        [(title, 46, WHITE, True)], align=PP_ALIGN.CENTER)
    box(s, Inches(0), Inches(2.0), EW, Inches(0.5),
        [(subtitle, 18, MUTE, False)], align=PP_ALIGN.CENTER)
    vw, vh = Inches(5.7), Inches(3.206)
    vx, vy = (EW - vw) / 2, Inches(2.72)
    vfr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, vx - Inches(0.06), vy - Inches(0.06),
                             vw + Inches(0.12), vh + Inches(0.12))
    vfr.fill.background(); vfr.line.color.rgb = GOLD; vfr.line.width = Pt(2.5)
    vfr.shadow.inherit = False
    mv = s.shapes.add_movie(mp4_path, vx, vy, vw, vh,
                            poster_frame_image=poster_path, mime_type="video/mp4")
    s.element.append(parse_xml(AUTOPLAY_TIMING.format(spid=mv.shape_id)))
    box(s, Inches(0), Inches(6.2), EW, Inches(0.55),
        [(footer, 17, GOLD, True)], align=PP_ALIGN.CENTER)
    if sign:
        box(s, Inches(0), Inches(6.82), EW, Inches(0.45),
            [(sign, 14, MUTE, False)], align=PP_ALIGN.CENTER)
    return s


def set_notes(prs, notes: dict):
    """notes = {頁碼(1-based): 逐字稿字串}。注意：插入待機頁會讓內容頁碼整體 +1。"""
    for i, sl in enumerate(prs.slides, 1):
        if i in notes:
            sl.notes_slide.notes_text_frame.text = notes[i]
